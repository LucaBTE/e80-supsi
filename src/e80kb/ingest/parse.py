"""File parsers: a file becomes chunks of text, each with its location."""

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

import docx
import openpyxl
import pypdf
from docx.table import Table
from pptx import Presentation

MAX_CHARS = 2000
SAMPLE_ROWS = 20
# First column of a requirement or test row: RQ-001, TST-003, TC-1.1.1
RECORD_ID = re.compile(r"^[A-Z]{2,4}-\d+(\.\d+)*$")


@dataclass
class Chunk:
    text: str
    loc: str  # page, slide, section or row
    record: dict[str, str] | None = None  # spreadsheet row whose first column is a record id


def parse(path: Path) -> list[Chunk]:
    return PARSERS[path.suffix.lower()](path)


def _pdf(path: Path) -> list[Chunk]:
    reader = pypdf.PdfReader(path)
    if reader.is_encrypted:
        reader.decrypt("")
    chunks = []
    for number, page in enumerate(reader.pages, 1):
        chunks += _split(page.extract_text() or "", f"p.{number}")
    return chunks


def _docx(path: Path) -> list[Chunk]:
    chunks, heading, lines = [], "body", []
    for block in docx.Document(path).iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = []
                previous = None
                for cell in row.cells:  # merged cells repeat
                    if cell._tc is not previous:
                        cells.append(cell.text.strip())
                    previous = cell._tc
                if any(cells):
                    lines.append(" | ".join(cells))
            continue
        style = (block.style.name if block.style is not None else "").lower()
        if style.startswith(("heading", "title")) and block.text.strip():
            chunks += _split("\n".join(lines), heading)
            heading, lines = block.text.strip(), []
        lines.append(block.text)
    return chunks + _split("\n".join(lines), heading)


def _pptx(path: Path) -> list[Chunk]:
    chunks = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        lines = [line for shape in slide.shapes for line in _shape_text(shape)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            lines.append(slide.notes_slide.notes_text_frame.text)
        chunks += _split("\n".join(lines), f"slide {number}")
    return chunks


def _shape_text(shape) -> list[str]:
    if hasattr(shape, "shapes"):  # group
        return [line for child in shape.shapes for line in _shape_text(child)]
    if shape.has_text_frame:
        return [shape.text_frame.text]
    if getattr(shape, "has_table", False):
        return [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
    return []


def _xlsx(path: Path) -> list[Chunk]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            chunk
            for sheet in workbook.worksheets
            for chunk in _table(sheet.title, [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)])
        ]
    finally:
        workbook.close()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    return str(value).strip()


def _csv(path: Path) -> list[Chunk]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _table(path.stem, [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)])


def _table(name: str, rows: list[list[str]]) -> list[Chunk]:
    """Requirement and test sheets give one chunk per row; other tables give a sample."""
    rows = [row for row in rows if any(row)]
    first = min((next(i for i, c in enumerate(row) if c) for row in rows), default=0)  # skip empty left columns
    rows = [row[first:] for row in rows]
    first_record = next((i for i, row in enumerate(rows) if row and RECORD_ID.fullmatch(row[0])), None)
    # A requirement sheet can have only an ID and a title column.
    minimum_columns = 2 if first_record is not None else 3
    header_at = next((i for i, row in enumerate(rows[:first_record])
                      if sum(1 for c in row if c) >= minimum_columns), None)
    if header_at is None:
        return _split("\n".join(" | ".join(c for c in row if c) for row in rows), name)

    header = [h or f"column {j}" for j, h in enumerate(rows[header_at], 1)]
    body = rows[header_at + 1:]
    records = [row for row in body if row and RECORD_ID.match(row[0])]
    if records:
        return [_record(name, header, row) for row in records]

    shown = body[:SAMPLE_ROWS]
    lines = [f"Table with {len(body)} rows, first {len(shown)} shown.", " | ".join(header)]
    lines += [" | ".join(row) for row in shown]
    return _split("\n".join(lines), f"{name} (sample)")


def _record(sheet: str, header: list[str], row: list[str]) -> Chunk:
    record = {}
    for h, value in zip(header, row):
        if value:
            record[h if h not in record else f"{h} ({len(record)})"] = value
    text = "\n".join(f"{h}: {v}" for h, v in record.items())
    return Chunk(text, f"{sheet}, row {row[0]}", record)


def _markdown(path: Path) -> list[Chunk]:
    chunks, heading, lines, in_code = [], "body", [], False
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if line.startswith("```"):
            in_code = not in_code
        elif not in_code and line.startswith("#") and line.lstrip("#").strip():
            chunks += _split("\n".join(lines), heading)
            heading, lines = line.lstrip("#").strip(), []
        lines.append(line)
    return chunks + _split("\n".join(lines), heading)


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self._skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _html(path: Path) -> list[Chunk]:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return _split("\n".join(parser.parts), "page")


def _split(text: str, loc: str) -> list[Chunk]:
    """Clean `text` and cut it into chunks of at most MAX_CHARS, on line breaks when possible."""
    text = "\n".join(line for line in (re.sub(r"\s+", " ", l).strip() for l in text.splitlines()) if line)
    pieces = []
    while len(text) > MAX_CHARS:
        cut = text.rfind("\n", 0, MAX_CHARS)
        cut = cut if cut > 0 else MAX_CHARS
        pieces.append(text[:cut])
        text = text[cut:].lstrip("\n")
    if text:
        pieces.append(text)
    if len(pieces) == 1:
        return [Chunk(pieces[0], loc)]
    return [Chunk(piece, f"{loc} (part {i})") for i, piece in enumerate(pieces, 1)]


PARSERS = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".csv": _csv,
    ".md": _markdown,
    ".txt": _markdown,
    ".html": _html,
    ".htm": _html,
}
