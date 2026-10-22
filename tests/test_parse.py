import openpyxl
from docx import Document

from e80kb.ingest.parse import MAX_CHARS, parse


def test_requirement_sheet_gives_one_chunk_per_row(tmp_path):
    path = tmp_path / "matrix.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Requirements Matrix"
    sheet.append(["H24 OFFERS RENEWAL - REQUIREMENTS MATRIX"])
    sheet.append(["Requirement ID", "Requirement Title", "Priority", "Requirement Status"])
    sheet.append(["Unique requirement identification", "Enter a concise title", "Level", "Status"])
    sheet.append(["RQ-001", "Automatic Contract Renewal Scheduling", "High", "Approved"])
    sheet.append(["RQ-002", "Avoid of Duplicate Renewal Process", "High", None])
    workbook.save(path)

    chunks = parse(path)

    assert [c.loc for c in chunks] == ["Requirements Matrix, row RQ-001", "Requirements Matrix, row RQ-002"]
    assert chunks[0].record == {
        "Requirement ID": "RQ-001",
        "Requirement Title": "Automatic Contract Renewal Scheduling",
        "Priority": "High",
        "Requirement Status": "Approved",
    }
    assert "Requirement Status" not in chunks[1].record


def test_data_table_gives_a_sample(tmp_path):
    path = tmp_path / "spedizioni.csv"
    path.write_text("data;importo;fornitore\n" + "".join(f"2026-01-{d:02d};{d}0;AZ{d}\n" for d in range(1, 31)))

    chunks = parse(path)

    assert len(chunks) == 1
    assert chunks[0].loc == "spedizioni (sample)"
    assert chunks[0].text.startswith("Table with 30 rows, first 20 shown.")
    assert chunks[0].record is None


def test_two_column_requirement_sheet(tmp_path):
    path = tmp_path / "requirements.csv"
    path.write_text("ID;Title\nRQ-001;First requirement\nRQ-002;Second requirement\n")

    chunks = parse(path)

    assert [c.record for c in chunks] == [
        {"ID": "RQ-001", "Title": "First requirement"},
        {"ID": "RQ-002", "Title": "Second requirement"},
    ]


def test_markdown_sections_and_long_text(tmp_path):
    path = tmp_path / "notes.md"
    long_line = "word " * 300
    path.write_text("# Intro\nhello\n```\n# not a heading\n```\n## Details\n" + "\n".join([long_line] * 3))

    chunks = parse(path)

    assert chunks[0].loc == "Intro"
    assert "# not a heading" in chunks[0].text
    assert [c.loc for c in chunks[1:]] == ["Details (part 1)", "Details (part 2)", "Details (part 3)"]
    assert all(len(c.text) <= MAX_CHARS for c in chunks)


def test_word_table_keeps_equal_values_in_distinct_cells(tmp_path):
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Merged heading"
    for cell, text in zip(table.rows[1].cells, ["Approved", "Approved", "Pending"]):
        cell.text = text
    document.save(path)

    text = "\n".join(chunk.text for chunk in parse(path))

    assert text.count("Merged heading") == 1
    assert "Approved | Approved | Pending" in text
